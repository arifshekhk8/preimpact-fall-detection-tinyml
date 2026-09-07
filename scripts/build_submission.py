#!/usr/bin/env python3
"""Assemble the CSC 471 submission.

    nothing/CSC471_Final_Assignment_SUBMISSION.docx
      1. Front page  -- the sample, retargeted to CSC 471, with a 5-row member table
      2. mp_paper.docx -- unchanged
      3. KPA justification form -- THE ORIGINAL TEMPLATE, filled in place
      4. Assignment rubric      -- THE ORIGINAL TEMPLATE, untouched, plus an
                                   evidence mapping appended after it

The two forms are opened and filled cell-by-cell rather than rebuilt, so their
layout, merges and styling are exactly as supplied. The rubric's "Marks Awarded"
column belongs to the instructor and is left blank.

The sample front page is for a different course (CSE 466, Career Planning); the
rubric and this assignment are CSC 471, so the front page is retargeted.

Run:  python scripts/build_submission.py
"""

from __future__ import annotations

import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docxcompose.composer import Composer

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "nothing"
OUT = SRC / "CSC471_Final_Assignment_SUBMISSION.docx"

COURSE_CODE = "CSC 471"
COURSE_NAME = "Microprocessor Based System"
INSTRUCTOR = "Dr. Md. Alomgir Hossain"
INSTRUCTOR_TITLE = "Course Instructor, Department of Computer Science and Engineering"
ASSIGNMENT = "Final Assignment"
SEMESTER = "Summer Semester 2026"
SECTION = "C"           # rubric covers C, D and E -- set yours
PROGRAM = "BCSE"
SUBMISSION_DATE = "08.09.2026"

# Only two members are known from the sample front page. The rest are placeholders
# and are marked so they cannot be missed.
MEMBERS = [
    ("Md. Arif Shekh", "23103022"),
    ("Md. Mahfuz Rana", "23103006"),
    ("<<MEMBER 3 NAME>>", "<<ID>>"),
    ("<<MEMBER 4 NAME>>", "<<ID>>"),
    ("<<MEMBER 5 NAME>>", "<<ID>>"),
]


# --------------------------------------------------------------------- helpers
def set_table_borders(table):
    """Apply single-line borders directly.

    Some of the supplied templates do not define a 'Table Grid' style, so relying on
    the named style raises. Writing the borders into the table properties works
    regardless of which styles a document happens to carry.
    """
    tblPr = table._tbl.tblPr
    for existing in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        borders.append(e)
    tblPr.append(borders)
    return table


def set_cell_bg(cell, hex_colour):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_colour)
    cell._tc.get_or_add_tcPr().append(shd)


def style_run(run, size=11, bold=False, font="Times New Roman"):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(a), font)
    return run


def para(doc, text="", size=11, bold=False, align=None,
         font="Times New Roman", space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        style_run(p.add_run(text), size, bold, font)
    return p


def cell_text(cell, text, size=10, bold=False, align=None, font="Times New Roman"):
    """Write into a cell, keeping its first paragraph (and thus the template's style)."""
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    style_run(p.add_run(text), size, bold, font)
    return p


def fill_blank_line(paragraph, label, value):
    """Replace a template's underscore rule with the value, keeping the label."""
    for r in list(paragraph.runs):
        r._element.getparent().remove(r._element)
    style_run(paragraph.add_run(f"{label}: "), 11, True)
    style_run(paragraph.add_run(value), 11, False)


# ============================================================== 1. FRONT PAGE
def build_front_page() -> Path:
    tmp = SRC / "_tmp_front.docx"
    shutil.copy(SRC / "front page.docx", tmp)
    doc = Document(tmp)

    # keep the university header + logo (paragraphs 0-2), rewrite everything below
    for p in doc.paragraphs[3:]:
        p._element.getparent().remove(p._element)

    para(doc, "", space_after=10)
    para(doc, ASSIGNMENT, size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_after=14)
    para(doc, f"Course Name: {COURSE_NAME}", size=15, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    para(doc, f"Course Code: {COURSE_CODE}", size=15, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    para(doc, f"Semester: {SEMESTER}", size=15, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)

    para(doc, "Submitted To", size=15, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    para(doc, INSTRUCTOR, size=15, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    para(doc, INSTRUCTOR_TITLE, size=12,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)

    para(doc, "Submitted By", size=15, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

    tbl = doc.add_table(rows=1 + len(MEMBERS), cols=3)
    set_table_borders(tbl)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(["No.", "Name", "Student ID"]):
        cell_text(tbl.rows[0].cells[j], h, size=12, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_bg(tbl.rows[0].cells[j], "D9D9D9")
    for i, (name, sid) in enumerate(MEMBERS, start=1):
        row = tbl.rows[i]
        cell_text(row.cells[0], str(i), size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(row.cells[1], name, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(row.cells[2], sid, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in tbl.rows:
        row.cells[0].width = Inches(0.6)
        row.cells[1].width = Inches(3.2)
        row.cells[2].width = Inches(1.8)

    para(doc, "", space_after=12)
    para(doc, f"Section: {SECTION}", size=14, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3)
    para(doc, f"Program: {PROGRAM}", size=14, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)
    para(doc, f"Date of Submission: {SUBMISSION_DATE}", size=14, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

    doc.save(tmp)
    return tmp


# ============================================ justification text (from the paper)
K_ROWS = {
    "K3": "Engineering fundamentals. Discrete-time signal processing is applied "
          "throughout: anti-alias filtering precedes decimation from 200 Hz to 50 Hz, "
          "because slicing without filtering aliases the impact transient into the "
          "passband. The same Nyquist reasoning excludes UP-Fall (18.4 Hz) and "
          "UMAFall's 20.1 Hz waist channel, since neither reaches 50 Hz without "
          "fabricating signal content. Datasets are also rotated onto one gravity "
          "convention using signed permutation matrices constrained to determinant +1, "
          "which preserves gyroscope handedness.",
    "K4": "Specialist knowledge. Full-integer INT8 quantisation, separable convolutions "
          "and global average pooling hold the network to 7,947 parameters and 22.5 KB. "
          "The work reports a non-obvious specialist result: instance normalisation "
          "placed inside the computation graph costs 30.95 points of macro-F1 under "
          "INT8, while the identical operation in the preprocessing pipeline costs "
          "nothing, because per-window mean and variance produce tensors whose dynamic "
          "range a single per-tensor quantisation scale cannot represent.",
    "K5": "Engineering tools. TensorFlow and Keras for training, TensorFlow Lite Micro "
          "for embedded inference, Kaggle T4 GPU sessions for the experiments, and the "
          "Arduino toolchain for the ESP32. Preprocessing is defined once in a "
          "version-controlled shared library imported by both the training scripts and "
          "the firmware generator, so the two cannot drift apart.",
    "K6": "Engineering practice. The system was built and validated on real hardware: an "
          "ESP32 DevKit V1 with an MPU6050 over I2C at 400 kHz, configured to ±16 g and "
          "±2000 dps, sampled by a hardware-timer interrupt at 50 Hz, driving a buzzer "
          "and status LED, with on-device instrumentation for inference latency and "
          "tensor-arena usage.",
    "K8": "Research literature. Four public datasets were audited against their source "
          "publications, and the results are positioned against the KFall benchmark (Yu "
          "et al., 2021), TinyFallNet (2023), PreFallKD (2023) and the cross-dataset "
          "study of Silva et al. (2024). Because that last work already establishes that "
          "a cross-dataset gap exists, the contribution here is deliberately narrowed to "
          "which mitigations are affordable on a microcontroller.",
}

P_ROWS = {
    "P2": "Addressed. This is the report's central engineering tension. Lead time and "
          "false-alarm rate conflict directly: requiring k consecutive windows before "
          "alarming cuts trial-level false positives from 538 to 26 out of 2,729 ADL "
          "trials, but each extra window costs one 0.5 s stride of lead time, and the "
          "mean lead of 663 ms is barely longer than a single stride. Model size, "
          "latency and accuracy conflict likewise: the proposed network gives up about "
          "two points of macro-F1 against the best comparator for a thirteen-fold "
          "reduction in parameters. Both trade-offs are quantified rather than asserted.",
    "P3": "Addressed. No textbook procedure answers the questions posed. The work "
          "required a tiered cross-validation design (5-fold subject-grouped for "
          "comparisons, full leave-one-subject-out for the headline result, and "
          "leave-one-dataset-out for the generalisation claim), an operating-point sweep "
          "over decision threshold and agreement count, and a null-controlled experiment "
          "to establish that the SisFall Enhanced annotations were unrecoverable — "
          "recovery peaked at 22.5 % against a time-reversed control and fell under "
          "harder filtering, which is what makes 22.5 % a ceiling rather than a waypoint.",
    "P4": "Addressed. Cross-dataset generalisation in wearable fall detection is an open "
          "research question rather than a routine exercise, and several problems "
          "encountered are not documented in the literature: that a dataset's gravity "
          "axis convention can masquerade as domain shift, and that the placement of "
          "instance normalisation relative to the quantisation boundary decides whether "
          "an INT8 model functions at all.",
    "P5": "Partially addressed. No single standard governs a research prototype of this "
          "kind, but the work follows established practice: IEEE conference format for "
          "reporting, I2C bus specifications for the sensor interface, TensorFlow Lite "
          "Micro conventions for embedded inference, and the reproducibility norms of the "
          "field — fixed random seeds, published preprocessing constants and fold-level "
          "result files released publicly.",
    "P6": "Addressed. Older adults at risk of falling are the end users, with caregivers "
          "and clinicians as secondary stakeholders whose tolerance for false alarms "
          "determines whether a device is worn at all. The BDT 1,450 bill of materials "
          "responds directly to affordability constraints in the Bangladeshi context, and "
          "the false-alarms-per-hour analysis exists precisely because a device that "
          "alarms hundreds of times an hour will be removed and never used, however "
          "accurate it is.",
    "P7": "Addressed. The subproblems cannot be solved independently. The sampling rate "
          "fixes the input tensor, which fixes model size and inference latency, which "
          "fixes the feasible stride, which bounds the achievable lead time and the "
          "false-alarm rate. The choice of normalisation strategy simultaneously "
          "determines cross-dataset transfer — it is the only adaptation method that "
          "helped — and whether the model survives quantisation. Changing any one of "
          "these forces the others to be revisited.",
}

A_ROWS = {
    "A1": "Addressed. Four public datasets totalling over 250,000 labelled windows; cloud "
          "GPU compute on Kaggle T4 accelerators; embedded hardware comprising ESP32, "
          "MPU6050, TP4056 charger and LiPo cell; software toolchains spanning "
          "TensorFlow, TensorFlow Lite Micro and Arduino; and the published research "
          "literature of the field.",
    "A2": "Addressed. The work required reconciling requirements across machine learning, "
          "embedded systems and signal processing, and coordination within a five-member "
          "group across data preparation, model training, firmware development, hardware "
          "assembly and report writing, using a shared version-controlled repository as "
          "the coordination mechanism.",
    "A3": "Addressed. Two findings appear not to have been reported previously: that "
          "instance normalisation must sit on the float side of the quantisation "
          "boundary, worth 30.95 points of macro-F1; and that the domain-adaptation "
          "ladder inverts, with the cheapest method the only one that improves transfer "
          "while CORAL and adversarial adaptation both degrade it. The argument that "
          "window-level specificity is close to meaningless for a continuously running "
          "detector is also not made elsewhere in this literature.",
    "A4": "Addressed. Falls are a leading cause of injury-related death among older "
          "adults, and a pre-impact alarm enables protective action rather than only "
          "post-hoc alerting. At BDT 1,450 the device is affordable in a low-resource "
          "setting. Running a 22.5 KB model on a microcontroller also avoids the energy "
          "cost and the privacy exposure of streaming continuous motion data to a "
          "server, since all inference happens on the device.",
    "A5": "Addressed. The work goes well beyond routine coursework: auditing public "
          "datasets against their source publications and finding four material defects, "
          "designing leakage-free evaluation protocols, diagnosing a silent quantisation "
          "failure before deployment, and running a neural network on a microcontroller "
          "with measured resource usage.",
}

RUBRIC_EVIDENCE = [
    ("CLO 1 — Microcontroller interfacing with sensors",
     "MPU6050 interfaced to the ESP32 over I2C at 400 kHz, configured by direct register "
     "writes to ±16 g (2048 LSB/g) and ±2000 dps (16.4 LSB/dps), with the on-chip "
     "low-pass filter enabled as an anti-alias stage and a sample-rate divider of 19 "
     "giving exactly 50 Hz. Acquisition runs from a hardware-timer interrupt into a "
     "100×6 ring buffer rather than a polled delay, because loop jitter accumulates and "
     "shifts window content away from the training distribution. Report Section V-E and "
     "Fig. 1(a)."),
    ("CLO 2 — Circuit design and implementation",
     "Complete prototype built and validated: ESP32 DevKit V1, MPU6050 on GPIO 21/22, "
     "active buzzer on GPIO 4, status LED on GPIO 2, powered from a 1000 mAh LiPo cell "
     "through a protected TP4056 charger. Components were selected against an explicit "
     "BDT 1,450 budget, and the ±16 g range was chosen deliberately because fall impacts "
     "exceed 8 g and a clipped peak destroys the signal the model depends on. "
     "Report Section V-E."),
    ("CLO 3 — Program development and logic flow",
     "Two documented, version-controlled codebases. The training library defines "
     "preprocessing once and is imported by both the training scripts and the firmware "
     "generator, so the two cannot diverge. The firmware implements the timer ISR, ring "
     "buffer, unit conversion, per-window instance normalisation, INT8 quantisation, "
     "inference, k-of-n agreement and cooldown, and additionally runs a rule-based "
     "detector on the same live samples for comparison. Report Fig. 1 and Section IV."),
    ("CLO 4 — Simulation/testing and result verification",
     "Verified at three levels. Offline: 5-fold subject-grouped cross-validation, full "
     "25-fold and 32-fold leave-one-subject-out, and leave-one-dataset-out across four "
     "corpora (Tables I–IV). Quantisation: FP32 against INT8 on identical held-out "
     "subjects, with 99.4 % agreement. On hardware: the firmware instruments its own "
     "inference latency and tensor-arena high-water mark. Every split is by subject or "
     "by dataset, never by window, to prevent leakage."),
    ("CLO 5 — Report writing and explanation of design",
     "Six-page report in IEEE conference format, with a method figure showing both the "
     "on-device pipeline and the network architecture (Fig. 1), a results figure "
     "quantifying the lead-time versus false-alarm trade-off (Fig. 2), and five tables. "
     "Limitations are stated explicitly rather than omitted, including simulated rather "
     "than real falls, an incomplete SisFall mirror, and the confounded UMAFall fold."),
]


# ==================================================== 3. KPA FORM (filled in place)
def build_kpa() -> Path:
    tmp = SRC / "_tmp_kpa.docx"
    shutil.copy(SRC / "KPA JUstification form.docx", tmp)
    doc = Document(tmp)

    names = ", ".join(n for n, _ in MEMBERS)
    ids = ", ".join(i for _, i in MEMBERS)
    fills = {
        "Course Code and Title": f"{COURSE_CODE} — {COURSE_NAME}",
        "Semester": SEMESTER,
        "Student Name": names,
        "Student ID": ids,
        "Course Instructor Name": INSTRUCTOR,
    }
    for p in doc.paragraphs:
        t = p.text.strip()
        for label, value in fills.items():
            if t.startswith(label) and "_" in t:
                fill_blank_line(p, label, value)
                break

    # Table 0: K rows fill column 2; P2-P7 rows fill the merged column 1
    t0 = doc.tables[0]
    for ri in range(1, 6):
        key = t0.rows[ri].cells[1].text.strip()
        if key in K_ROWS:
            cell_text(t0.rows[ri].cells[2], K_ROWS[key], size=9,
                      align=WD_ALIGN_PARAGRAPH.LEFT)
    for row in t0.rows[6:]:
        label = row.cells[0].text.strip()
        m = re.match(r"(P[2-7])", label)
        if m and m.group(1) in P_ROWS:
            cell_text(row.cells[1], P_ROWS[m.group(1)], size=9,
                      align=WD_ALIGN_PARAGRAPH.LEFT)

    # Table 1: A rows
    t1 = doc.tables[1]
    for row in t1.rows[1:]:
        m = re.match(r"(A[1-5])", row.cells[0].text.strip())
        if m and m.group(1) in A_ROWS:
            cell_text(row.cells[1], A_ROWS[m.group(1)], size=9,
                      align=WD_ALIGN_PARAGRAPH.LEFT)

    doc.save(tmp)
    return tmp


# ============================== 4. RUBRIC (template untouched + evidence appended)
def build_rubric() -> Path:
    tmp = SRC / "_tmp_rubric.docx"
    shutil.copy(SRC / "Assignment Rubrics for csc 471.docx", tmp)
    doc = Document(tmp)

    doc.add_page_break()
    para(doc, "Evidence Against Each Criterion", size=13, bold=True, space_after=4)
    para(doc, "The 'Marks Awarded' column in the rubric above is left blank for the "
              "course instructor. This table records where in the report each assessment "
              "criterion is demonstrated.", size=10, space_after=8)

    t = doc.add_table(rows=1 + len(RUBRIC_EVIDENCE), cols=2)
    set_table_borders(t)
    cell_text(t.rows[0].cells[0], "Course Learning Outcome", size=10, bold=True)
    cell_text(t.rows[0].cells[1], "How it is met in this submission", size=10, bold=True)
    for c in t.rows[0].cells:
        set_cell_bg(c, "D9D9D9")
    for r, (clo, ev) in enumerate(RUBRIC_EVIDENCE, start=1):
        cell_text(t.rows[r].cells[0], clo, size=9, bold=True)
        cell_text(t.rows[r].cells[1], ev, size=9)
        t.rows[r].cells[0].width = Inches(1.9)
        t.rows[r].cells[1].width = Inches(5.0)

    doc.save(tmp)
    return tmp



# ================================================================= PDF export
# Word is not scriptable here, so LibreOffice does the conversion. It renders the
# paper's figures and both filled forms faithfully -- verified against the source.
SOFFICE_CANDIDATES = (
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "soffice",
    "libreoffice",
)


def export_pdf(docx_path: Path) -> Path | None:
    """Convert the assembled document to PDF next to it."""
    exe = next((c for c in SOFFICE_CANDIDATES
                if Path(c).exists() or shutil.which(c)), None)
    if exe is None:
        print("  (no LibreOffice found -- skipping PDF; open the .docx and "
              "'Save as PDF' by hand)")
        return None

    proc = subprocess.run(
        [exe, "--headless", "--convert-to", "pdf",
         "--outdir", str(docx_path.parent), str(docx_path)],
        capture_output=True, text=True, timeout=600,
    )
    pdf = docx_path.with_suffix(".pdf")
    if proc.returncode != 0 or not pdf.exists():
        print(f"  !! PDF conversion failed: {proc.stdout}{proc.stderr}")
        return None
    return pdf


# ==================================================================== assemble
def main() -> int:
    front, kpa, rubric = build_front_page(), build_kpa(), build_rubric()

    base = Document(front)
    base.add_page_break()
    comp = Composer(base)
    comp.append(Document(SRC / "mp_paper.docx"))
    comp.append(Document(kpa))
    comp.append(Document(rubric))
    comp.save(OUT)

    for t in (front, kpa, rubric):
        t.unlink(missing_ok=True)

    d = Document(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")
    print(f"  paragraphs {len(d.paragraphs)}  tables {len(d.tables)}  "
          f"images {len(d.inline_shapes)}")
    blanks = sum(1 for t in d.tables for r in t.rows for c in r.cells
                 if not c.text.strip())
    print(f"  empty table cells remaining: {blanks}")

    pdf = export_pdf(OUT)
    if pdf:
        try:
            from pypdf import PdfReader
            print(f"wrote {pdf.relative_to(ROOT)}  ({len(PdfReader(pdf).pages)} pages)")
        except ImportError:
            print(f"wrote {pdf.relative_to(ROOT)}")
    todo = [n for n, _ in MEMBERS if n.startswith("<<")]
    if todo:
        print(f"\n  !! {len(todo)} placeholder member rows -- search for '<<'")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
